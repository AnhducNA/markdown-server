"""Integration tests for FastAPI endpoints using TestClient."""
import io
import json
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

from app.main import app

client = TestClient(app)


def test_root():
    resp = client.get("/")
    assert resp.status_code == 200
    data = resp.json()
    assert "docs" in data
    assert "health" in data


def test_health():
    resp = client.get("/api/v1/health")
    assert resp.status_code == 200
    data = resp.json()
    assert data["status"] == "ok"
    assert "available_parsers" in data
    assert "pymupdf" in data["available_parsers"]


def test_convert_invalid_extension():
    """Upload a .txt file should return 400."""
    fake_txt = io.BytesIO(b"just some text")
    resp = client.post(
        "/api/v1/convert",
        files={"file": ("document.txt", fake_txt, "text/plain")},
    )
    assert resp.status_code == 400
    assert "Unsupported file format" in resp.json()["detail"]


def test_convert_url_invalid():
    """POST with invalid URL should return validation error."""
    resp = client.post(
        "/api/v1/convert-url",
        json={"url": "not-a-url"},
    )
    assert resp.status_code == 422


def test_document_not_found():
    resp = client.get("/api/v1/documents/nonexistent_doc_999")
    assert resp.status_code == 404


def test_markdown_not_found():
    resp = client.get("/api/v1/documents/nonexistent_doc_999/markdown")
    assert resp.status_code == 404


def test_manifest_not_found():
    resp = client.get("/api/v1/documents/nonexistent_doc_999/manifest")
    assert resp.status_code == 404


def test_convert_docx(tmp_path: Path):
    """Upload a minimal valid DOCX and verify conversion pipeline."""
    import docx

    # Build a minimal DOCX in memory
    document = docx.Document()
    document.add_heading("Test Document", level=1)
    document.add_paragraph("This is a test paragraph.")
    document.add_heading("Section 1", level=2)
    document.add_paragraph("Content of section 1.")

    docx_path = tmp_path / "test.docx"
    document.save(str(docx_path))

    with open(docx_path, "rb") as f:
        resp = client.post(
            "/api/v1/convert",
            files={"file": ("test.docx", f, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            data={"document_id": "api_test_docx"},
        )

    assert resp.status_code == 200, f"Expected 200, got {resp.status_code}: {resp.text}"
    data = resp.json()
    assert data["status"] == "completed"
    assert "api_test_docx" in data["document_id"]
    assert data["assets_count"] >= 0

    # Fetch the generated markdown
    md_resp = client.get(f"/api/v1/documents/{data['document_id']}/markdown")
    assert md_resp.status_code == 200
    md_text = md_resp.text
    assert "Test Document" in md_text

    # Fetch manifest
    manifest_resp = client.get(f"/api/v1/documents/{data['document_id']}/manifest")
    assert manifest_resp.status_code == 200
    manifest = manifest_resp.json()
    assert manifest["source_type"] == "docx"
    assert manifest["blocks_count"] > 0
    assert len(manifest["chunks"]) > 0
